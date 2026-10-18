import io
import re
import tempfile
import zipfile
from datetime import date, datetime
from pathlib import Path

from django.core.files.base import ContentFile, File
from django.db import IntegrityError, transaction

from apps.declarations import services as decl_services
from apps.declarations.models import Attachment, Declaration, Payment
from apps.registry.models import CitesDocument, Specimen
from apps.species.models import Species
from apps.vault.storage import read_plain

from .extract import IMAGE_EXT, PDF_EXT, SHEET_EXT, extract
from .models import ImportedFile, file_sha256
from .parse import parse

ALLOWED_EXT = PDF_EXT | IMAGE_EXT | SHEET_EXT | {".docx", ".odt", ".odg"}


def _year_hint(path: str) -> int | None:
    m = re.search(r"(?:^|[/\\])((?:19|20)\d{2})(?:[/\\]|$)", path) or re.search(r"((?:19|20)\d{2})", path)
    return int(m.group(1)) if m else None


def add_file(batch, django_file, original_name: str, original_path: str = "") -> tuple[ImportedFile | None, str]:
    """Store a file in the batch (deduplicated per keeper by SHA-256) and analyze it."""
    ext = Path(original_name).suffix.lower()
    if ext not in ALLOWED_EXT:
        return None, f"{original_name}: formato non supportato"
    sha = file_sha256(django_file, batch.keeper)
    existing = ImportedFile.objects.filter(keeper=batch.keeper, sha256=sha).first()
    if existing:
        return None, f"{original_name}: già importato ({existing.original_path or existing.original_name})"
    imported = ImportedFile(
        batch=batch, keeper=batch.keeper, original_name=original_name, original_path=original_path or original_name,
        sha256=sha, year_hint=_year_hint(original_path or original_name),
    )
    imported.file.save(original_name, django_file, save=False)
    try:
        imported.save()
    except IntegrityError:
        return None, f"{original_name}: già importato"
    analyze(imported)
    return imported, ""


def add_zip(batch, uploaded) -> list[str]:
    msgs = []
    with zipfile.ZipFile(uploaded) as zf:
        for info in zf.infolist():
            if info.is_dir() or "__MACOSX" in info.filename or Path(info.filename).name.startswith("."):
                continue
            data = zf.read(info)
            name = Path(info.filename).name
            _, msg = add_file(batch, ContentFile(data, name=name), name, info.filename)
            if msg:
                msgs.append(msg)
    return msgs


def add_path(batch, root: Path) -> list[str]:
    msgs = []
    for path in sorted(p for p in Path(root).rglob("*") if p.is_file()):
        with path.open("rb") as fh:
            _, msg = add_file(batch, File(fh, name=path.name), path.name, str(path.relative_to(root)))
        if msg:
            msgs.append(msg)
    return msgs


def analyze(imported: ImportedFile) -> ImportedFile:
    ext = Path(imported.original_name).suffix.lower()
    known = list(Species.objects.values_list("scientific_name", flat=True))
    try:
        if ext in SHEET_EXT:
            imported.detected_kind = "SHEET"
            imported.confidence = 0.9
            imported.method = "xlsx"
            imported.parsed = {"kind": "SHEET", "data": {}, "rows": [], "notes": []}
        else:
            # stored file is encrypted: extract from a decrypted temporary copy
            with tempfile.TemporaryDirectory() as tmp:
                plain = Path(tmp) / f"doc{Path(imported.original_name).suffix.lower()}"
                plain.write_bytes(read_plain(imported.file))
                ex = extract(plain)
            result = parse(imported.original_name, ex.text, ex.form_fields, known)
            imported.text = ex.text[:50000]
            imported.method = ex.method
            imported.detected_kind = result.kind
            imported.confidence = result.confidence
            imported.parsed = result.as_dict()
            imported.parsed["notes"] = ex.warnings + result.notes
        imported.status = ImportedFile.Status.PARSED
        imported.error = ""
    except Exception as exc:  # keep the file, let the user enter data by hand
        imported.status = ImportedFile.Status.ERROR
        imported.error = str(exc)
    imported.save()
    return imported


def _attach(declaration: Declaration, imported: ImportedFile, kind=Attachment.Kind.ORIGINAL):
    imported.file.open("rb")
    try:
        att = Attachment(declaration=declaration, kind=kind, description=imported.original_name)
        att.file.save(imported.original_name, File(imported.file, name=imported.original_name), save=True)
    finally:
        imported.file.close()
    return att


def _status_fields(declaration, cd):
    declaration.imported = True
    declaration.deadline = None
    declaration.signed_on = cd.get("signed_on") or declaration.event_date
    declaration.sent_on = cd.get("signed_on") or declaration.event_date
    if cd.get("protocol_number"):
        declaration.protocol_number = cd["protocol_number"]
        declaration.protocol_date = cd.get("protocol_date")
        declaration.status = Declaration.Status.PROTOCOLLED
    else:
        declaration.status = Declaration.Status.SENT
    declaration.save()


@transaction.atomic
def apply(imported: ImportedFile, cd: dict, user) -> str:
    """Turn a reviewed import into registry records. Returns a human message."""
    keeper = imported.keeper
    action = cd["action"]

    if action == "skip":
        imported.status = ImportedFile.Status.SKIPPED
        imported.save()
        return "File ignorato."

    if action == "birth":
        rows = [decl_services.NewbornSpec(quantity=cd["quantity"] or 1)]
        result = decl_services.register_birth(
            keeper=keeper, user=user, species=cd["species"], birth_date=cd["event_date"], rows=rows,
            breeding_group=cd.get("breeding_group") or "",
        )
        decl = result.declaration
        if decl is None:  # species not in Annex A/B: no form, keep the record anyway
            decl = Declaration.objects.create(keeper=keeper, kind=Declaration.Kind.BIRTH_B, event_date=cd["event_date"],
                                              imported=True, created_by=user)
        _status_fields(decl, cd)
        _attach(decl, imported)
        if cd.get("protocol_number"):
            decl_services.mark_protocolled(decl, cd["protocol_number"], cd.get("protocol_date"))
        imported.declaration = decl
        msg = f"Nascita registrata: {cd['quantity']} × {cd['species']}."

    elif action == "death":
        specimen = cd.get("specimen")
        if specimen:
            decl = decl_services.register_death(
                keeper=keeper, user=user, specimen=specimen, quantity=cd["quantity"] or 1,
                death_date=cd["event_date"], cause=cd.get("notes") or "",
            )
        else:
            decl = None
        if decl is None:
            decl = Declaration.objects.create(
                keeper=keeper, kind=Declaration.Kind.DEATH, event_date=cd["event_date"], imported=True,
                created_by=user, notes=f"Importata senza esemplare collegato: {cd['quantity']} × {cd['species']}",
            )
        _status_fields(decl, cd)
        _attach(decl, imported)
        if cd.get("protocol_number"):
            decl_services.mark_protocolled(decl, cd["protocol_number"], cd.get("protocol_date"))
        imported.declaration = decl
        msg = "Decesso registrato."

    elif action == "document":
        doc = CitesDocument.objects.create(
            keeper=keeper, doc_type=cd["doc_type"], number=cd["number"], issued_on=cd.get("issued_on"),
            issuer=str(keeper.cites_office or ""), notes=cd.get("notes") or "",
        )
        imported.file.open("rb")
        try:
            doc.file.save(imported.original_name, File(imported.file, name=imported.original_name), save=True)
        finally:
            imported.file.close()
        if cd.get("specimens"):
            doc.specimens.set(cd["specimens"])
        link = cd.get("link_declaration")
        if link:
            link.protocol_number = link.protocol_number or cd["number"]
            link.protocol_date = link.protocol_date or cd.get("issued_on")
            if link.status in (Declaration.Status.DRAFT, Declaration.Status.GENERATED, Declaration.Status.SENT):
                link.status = Declaration.Status.PROTOCOLLED
            link.save()
            _attach(link, imported, Attachment.Kind.REPLY)
            if not cd.get("specimens"):
                doc.specimens.set([i.specimen for i in link.items.all()])
            imported.declaration = link
        imported.cites_document = doc
        msg = f"Documento {doc.get_doc_type_display()} {doc.number} registrato."

    elif action == "payment":
        pay = Payment.objects.create(
            keeper=keeper, amount=cd["amount"], paid_on=cd["paid_on"], reference=cd.get("number") or "",
            description=cd.get("notes") or imported.original_name, fee_type=cd.get("fee_type"),
            declaration=cd.get("link_declaration"),
        )
        imported.file.open("rb")
        try:
            pay.receipt.save(imported.original_name, File(imported.file, name=imported.original_name), save=True)
        finally:
            imported.file.close()
        msg = "Pagamento registrato."

    elif action == "archive":
        link = cd.get("link_declaration")
        if link:
            _attach(link, imported, Attachment.Kind.OTHER)
            imported.declaration = link
        imported.status = ImportedFile.Status.ARCHIVED
        imported.save()
        return "File archiviato."
    else:
        raise ValueError(action)

    imported.status = ImportedFile.Status.DONE
    imported.save()
    return msg


# --------------------------------------------------------------------------- spreadsheet import

SHEET_COLUMNS = ["specie", "quantita", "sesso", "data_nascita", "nome", "tipo_marcatura", "marcatura",
                 "documento", "data_documento", "provenienza", "note"]


def sheet_template() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "esemplari"
    ws.append(SHEET_COLUMNS)
    ws.append(["Testudo hermanni", 1, "F", "15/06/2015", "Femmina 1", "microchip", "380260000000000",
               "PROT 123 DEL 1999", "", "pre", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_date(value):
    from .parse import parse_date

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return parse_date(str(value)) if value else None


@transaction.atomic
def import_specimen_sheet(keeper, user, fileobj) -> tuple[int, list[str]]:
    """Bulk-load existing animals (consistenza iniziale) from an XLSX with SHEET_COLUMNS headers."""
    from openpyxl import load_workbook

    wb = load_workbook(fileobj, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return 0, ["Foglio vuoto."]
    headers = [str(h or "").strip().lower() for h in rows[0]]
    errors, count = [], 0
    for n, raw in enumerate(rows[1:], start=2):
        rec = dict(zip(headers, raw))
        if not any(rec.values()):
            continue
        name = str(rec.get("specie") or "").strip()
        species = Species.objects.filter(scientific_name__iexact=name).first()
        if not species:
            errors.append(f"Riga {n}: specie '{name}' sconosciuta (aggiungerla prima).")
            continue
        doc = None
        if rec.get("documento"):
            doc = CitesDocument.objects.create(
                keeper=keeper, doc_type=CitesDocument.DocType.OLD_DECLARATION, number=str(rec["documento"]),
                issued_on=_cell_date(rec.get("data_documento")),
            )
        kind = str(rec.get("tipo_marcatura") or "").strip().lower()
        kind = {"microchip": "microchip", "anello": "ring", "ring": "ring", "foto": "photo"}.get(kind, "other" if kind else "")
        origin = str(rec.get("provenienza") or "pre").strip().lower()
        origin = origin if origin in Specimen.Origin.values else Specimen.Origin.PRE_EXISTING
        sex = str(rec.get("sesso") or "I").strip().upper()[:1]
        try:
            with transaction.atomic():
                decl_services.register_acquisition(
                    keeper=keeper, user=user, species=species, quantity=int(rec.get("quantita") or 1),
                    acquired_on=_cell_date(rec.get("data_nascita")) or _cell_date(rec.get("data_documento"))
                    or date.today(),
                    counterparty=None, sex=sex if sex in "MFI" else "I", birth_date=_cell_date(rec.get("data_nascita")),
                    name=str(rec.get("nome") or ""), marking_kind=kind, marking_code=str(rec.get("marcatura") or ""),
                    document=doc, origin=origin,
                )
            count += 1
        except Exception as exc:
            errors.append(f"Riga {n}: {exc}")
    return count, errors
